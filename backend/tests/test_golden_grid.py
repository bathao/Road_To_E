"""Golden (characterization) tests for the Daily Tracker grid and its export.

``_grid_cells`` is the one renderer behind GET /weeks, the CSV and the XLSX
export; it is also the riskiest piece of the planned tracker/service split
(REFACTOR.md §3.1). These tests pin its output for one busy week in the past
so a refactor that moves code around cannot change a single cell unnoticed.

The week is 2026-06-01 (Mon) → 2026-06-07 and lies entirely before the
default ELO anchor (2026-07-27), so every match is "before anchor" and no
test ever depends on the wall clock: every day is in the past, hence the
empty Sunday renders red.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import json

from openpyxl import load_workbook

from conftest import assert_golden, category_id
from app.features.tournament import schemas as t_schemas
from app.features.tournament import service as t_service
from app.features.tracker import service
from app.features.tracker.models import Activity, DayNote, PhysicalCheck, Player
from app.features.training import service as training_service

WEEK_START = dt.date(2026, 6, 1)  # a Monday
WEEK_END = dt.date(2026, 6, 7)


def _d(offset: int) -> str:
    return (WEEK_START + dt.timedelta(days=offset)).isoformat()


def _seed_week(client, db) -> None:
    coach_cat = category_id(db, "train_with_coach")
    partner_cat = category_id(db, "training_with_partner")
    serve_cat = category_id(db, "serve")
    practice = category_id(db, "practice_match")
    official = category_id(db, "official_match")
    tournament = category_id(db, "tournament_match")
    coaches = {c.name: c for c in service.list_coaches(db)}
    minh, phi = coaches["Minh Thới"], coaches["Phi Vũ"]

    anna = Player(name="Anna", points=950)
    binh = Player(name="Bình", points=1200, plays_pips=True)
    cuong = Player(name="Cường", points=700)
    buddy = Player(name="Đồng Đội", points=900)
    db.add_all([anna, binh, cuong, buddy])
    db.commit()

    db.add_all([
        # Mon: package coach, first session of a new 10-session block.
        Activity(date=WEEK_START, category_id=coach_cat, duration_minutes=60,
                 coach_id=minh.id, is_package_start=True),
        # Tue: per-session coach whose legacy note repeats his name (the
        # stutter guard must swallow it).
        Activity(date=WEEK_START + dt.timedelta(days=1), category_id=coach_cat,
                 duration_minutes=90, coach_id=phi.id, note="Phi Vũ"),
        # Wed: per-session coach with a real note.
        Activity(date=WEEK_START + dt.timedelta(days=2), category_id=coach_cat,
                 duration_minutes=60, coach_id=phi.id, note="footwork"),
        Activity(date=WEEK_START, category_id=partner_cat, duration_minutes=45),
        Activity(date=WEEK_START + dt.timedelta(days=3), category_id=serve_cat,
                 duration_minutes=20),
        # Legacy physical checklist (before the Training Center cutover).
        *[PhysicalCheck(date=WEEK_START, item_key=k)
          for k in ("wall_sit", "plank", "squats", "stretching", "sit_ups")],
        PhysicalCheck(date=WEEK_START + dt.timedelta(days=1), item_key="plank"),
        DayNote(date=WEEK_START,
                text="Knee felt fine today; long rallies in the evening, "
                     "need to keep the elbow lower on the backhand."),
        DayNote(date=WEEK_END, text="Rest day"),
    ])
    db.commit()

    # Training Center takes over the Physical row from Thursday on.
    state = training_service.ensure_state(db)
    state.cutover_date = WEEK_START + dt.timedelta(days=3)
    db.commit()
    s1 = training_service._materialise(db, "foundation", 1)
    items1 = sorted(s1.items, key=lambda it: it.sort_order)
    for it in items1[:2]:
        training_service.tick_item(db, s1.id, it.id, True)
    training_service.complete_session(
        db, "foundation", 1, "ok", done_on=WEEK_START + dt.timedelta(days=4)
    )
    s2 = training_service._materialise(db, "foundation", 2)
    for it in s2.items:
        training_service.tick_item(db, s2.id, it.id, True)
    training_service.complete_session(
        db, "foundation", 2, None, done_on=WEEK_START + dt.timedelta(days=5)
    )

    entry_id = t_service.create_tournament(db, t_schemas.TournamentIn(
        name="Gia Định Open",
        start_date=WEEK_START + dt.timedelta(days=3),
        entries=[t_schemas.EntryIn(discipline="singles")],
    )).tournaments[0].entries[0].id

    matches = [
        {"date": _d(0), "category_id": practice, "discipline": "singles",
         "my_sets": 3, "opp_sets": 1, "opponent_id": anna.id},
        {"date": _d(0), "category_id": practice, "discipline": "singles",
         "my_sets": 1, "opp_sets": 3, "opponent_id": binh.id,
         "handicap": -1, "handicap_pattern": "2-0-2"},
        {"date": _d(2), "category_id": official, "discipline": "doubles",
         "my_sets": 3, "opp_sets": 2, "partner_id": buddy.id,
         "opponent_id": anna.id, "opponent2_id": cuong.id},
        {"date": _d(3), "category_id": tournament, "discipline": "singles",
         "my_sets": 2, "opp_sets": 3, "opponent_id": cuong.id,
         "tournament_entry_id": entry_id, "round": "qf"},
        {"date": _d(4), "category_id": official, "discipline": "one_v_two",
         "best_of": 3, "my_sets": 2, "opp_sets": 0,
         "opponent_id": anna.id, "opponent2_id": cuong.id, "handicap": 2},
        {"date": _d(4), "category_id": official, "discipline": "two_v_one",
         "my_sets": 1, "opp_sets": 3, "partner_id": buddy.id,
         "opponent_id": binh.id, "event_name": "Thứ sáu CLB"},
        {"date": _d(5), "category_id": practice, "is_nonplaying": True,
         "nonplaying_label": "Travel"},
    ]
    for payload in matches:
        r = client.post("/api/tracker/matches", json=payload)
        assert r.status_code == 200, r.text


def test_week_grid_golden(client, db):
    _seed_week(client, db)
    week = client.get("/api/tracker/weeks", params={"start": WEEK_START.isoformat()})
    assert week.status_code == 200
    body = week.json()
    assert_golden(
        "week_2026-06-01.json",
        json.dumps(body, ensure_ascii=False, indent=1, sort_keys=True) + "\n",
    )

    # A narrower range renders the same cells for the days it covers (the
    # Overall colour of a day never depends on the range around it).
    one_day = client.get(
        "/api/tracker/weeks", params={"start": _d(2), "end": _d(2)}
    ).json()
    assert one_day["days"] == [_d(2)]
    for key, cell in one_day["cells"].items():
        assert body["cells"][key] == cell, key


def test_week_export_golden(client, db):
    _seed_week(client, db)
    rng = {"from": WEEK_START.isoformat(), "to": WEEK_END.isoformat()}

    csv_res = client.get("/api/tracker/export", params={**rng, "format": "csv"})
    assert csv_res.status_code == 200
    assert csv_res.headers["content-disposition"].endswith('tracker_2026-06-01_2026-06-07.csv"')
    csv_text = csv_res.content.decode("utf-8-sig")
    assert_golden("week_2026-06-01.csv", csv_text)

    # The XLSX is the same grid: every text cell equals the CSV cell (the CSV
    # flattens newlines to " | "), and the Overall row carries the colour the
    # CSV spells out as a word.
    xlsx_res = client.get("/api/tracker/export", params={**rng, "format": "xlsx"})
    assert xlsx_res.status_code == 200
    ws = load_workbook(io.BytesIO(xlsx_res.content)).active
    csv_rows = list(csv.reader(io.StringIO(csv_text)))
    header = [c.value for c in ws[1]]
    assert header[0] == "Category" and len(header) == len(csv_rows[0])
    for i, csv_row in enumerate(csv_rows[1:], start=2):
        label = ws.cell(row=i, column=1).value
        assert label == csv_row[0]
        for j in range(2, len(csv_row) + 1):
            cell = ws.cell(row=i, column=j)
            expected = csv_row[j - 1]
            if label == "Overall":
                assert cell.value is None
                fill = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None
                assert fill == service._RATING_HEX.get(expected), (label, j, expected)
            else:
                got = (cell.value or "").replace("\n", " | ")
                assert got == expected, (label, j)
