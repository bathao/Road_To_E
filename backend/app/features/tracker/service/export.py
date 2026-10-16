"""CSV / XLSX export of the grid.

Part of the tracker service package (split 2026-10-16 from the former
2.7k-line service.py; see REFACTOR.md §3.1). Import names via
``app.features.tracker.service`` — the package __init__ re-exports them.
"""
from __future__ import annotations

import csv
import datetime as dt
import io

from sqlalchemy.orm import Session

from .grid import _grid_cells, _load_range
from .periods import _date_range

_RATING_HEX = {"green": "FF63BE7B", "yellow": "FFFFEB84", "red": "FFE06666"}


_COLOR_HEX = {"green": "FF92D050", "yellow": "FFFFFF00", "none": None}


def _build_grid(db: Session, date_from: dt.date, date_to: dt.date):
    """Return (categories, dates, cell_text, rating_by_date, cell_colors).

    Same _grid_cells renderer as the on-screen grid, so the export can never
    silently drop information the grid shows."""
    dates = _date_range(date_from, date_to)
    rng = _load_range(db, date_from, date_to)
    text, cell_colors, rating_by_date = _grid_cells(db, rng, dates, for_export=True)
    return rng.categories, dates, text, rating_by_date, cell_colors


def export_csv(db: Session, date_from: dt.date, date_to: dt.date) -> bytes:
    categories, dates, text, rating_by_date, _colors = _build_grid(
        db, date_from, date_to
    )
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Category", *[d.isoformat() for d in dates]])
    for c in categories:
        row = [c.label]
        for d in dates:
            iso = d.isoformat()
            if c.key == "overall":
                row.append(rating_by_date.get(iso, ""))
            else:
                row.append(text.get((c.id, iso), "").replace("\n", " | "))
        writer.writerow(row)
    return buf.getvalue().encode("utf-8-sig")


def export_xlsx(db: Session, date_from: dt.date, date_to: dt.date) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    categories, dates, text, rating_by_date, cell_colors = _build_grid(
        db, date_from, date_to
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Daily Tracker"

    header_font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    ws.cell(row=1, column=1, value="Category").font = header_font
    for j, d in enumerate(dates, start=2):
        cell = ws.cell(row=1, column=j, value=d.strftime("%a %d %b %Y"))
        cell.font = header_font

    for i, c in enumerate(categories, start=2):
        label_cell = ws.cell(row=i, column=1, value=c.label)
        label_cell.font = header_font
        hex_color = _COLOR_HEX.get(c.color_group)
        if hex_color:
            label_cell.fill = PatternFill("solid", fgColor=hex_color)
        for j, d in enumerate(dates, start=2):
            iso = d.isoformat()
            cell = ws.cell(row=i, column=j)
            if c.key == "overall":
                overall_color = rating_by_date.get(iso)
                if overall_color and overall_color in _RATING_HEX:
                    cell.fill = PatternFill("solid", fgColor=_RATING_HEX[overall_color])
            else:
                cell.value = text.get((c.id, iso), "")
                cell.alignment = wrap
                fill = cell_colors.get((c.id, iso))
                if fill and fill in _RATING_HEX:
                    cell.fill = PatternFill("solid", fgColor=_RATING_HEX[fill])

    ws.column_dimensions["A"].width = 26
    for col in range(2, len(dates) + 2):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 18

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
